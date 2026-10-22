"""Employee registry: create, import, duplicates, invite binding, own consent."""

from __future__ import annotations

import csv
import io
import secrets
from dataclasses import dataclass
from datetime import date, datetime
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.domain.rules import duplicate_reasons, normalize_username
from app.errors import DomainError, DuplicateEmployeeError, NotFoundError
from app.models import Department, Employee, utcnow
from app.services.audit import write_audit

TRUE_VALUES = {"1", "true", "yes", "да", "y"}


@dataclass
class EmployeeInput:
    full_name: str
    preferred_name: str
    birth_date: date
    department_id: UUID | None = None
    position: str = ""
    phone: str = ""
    city: str = ""
    telegram_id: int | None = None
    telegram_username: str | None = None
    hired_at: date | None = None
    public_greeting_consent: bool = False
    participate_in_collections: bool = True
    notification_settings: dict | None = None
    status: str = "active"
    bot_status: str | None = None


def create_employee(session: Session, data: EmployeeInput, actor_id: UUID | None) -> Employee:
    _validate(data)
    matches = _matches(session, data, exclude_id=None, source="database")
    if matches:
        raise DuplicateEmployeeError(matches)
    employee = _build(data)
    session.add(employee)
    write_audit(session, actor_id, "employee.create", "employee", employee.id, {"full_name": employee.full_name})
    session.commit()
    return employee


def update_employee(session: Session, employee_id: UUID, data: EmployeeInput, actor_id: UUID | None) -> Employee:
    employee = _get(session, employee_id)
    _validate(data)
    matches = _matches(session, data, exclude_id=employee_id, source="database")
    if matches:
        raise DuplicateEmployeeError(matches)
    employee.full_name = data.full_name.strip()
    employee.preferred_name = data.preferred_name.strip()
    employee.birth_date = data.birth_date
    employee.department_id = data.department_id
    employee.position = data.position.strip()
    employee.phone = (data.phone or "").strip()
    employee.city = (data.city or "").strip()
    employee.telegram_id = data.telegram_id
    employee.telegram_username = normalize_username(data.telegram_username)
    employee.hired_at = data.hired_at
    employee.public_greeting_consent = data.public_greeting_consent
    employee.participate_in_collections = data.participate_in_collections
    if data.notification_settings is not None:
        employee.notification_settings = {
            "announcements": bool(data.notification_settings.get("announcements", True)),
            "collections": bool(data.notification_settings.get("collections", True)),
        }
    if data.status in {"active", "inactive"}:
        employee.status = data.status
    if employee.telegram_id and employee.bot_status == "not_started":
        employee.bot_status = "active"
    employee.updated_at = utcnow()
    write_audit(session, actor_id, "employee.update", "employee", employee.id, {})
    session.commit()
    return employee


def deactivate_employee(session: Session, employee_id: UUID, actor_id: UUID | None) -> Employee:
    employee = _get(session, employee_id)
    employee.status = "inactive"
    employee.updated_at = utcnow()
    write_audit(session, actor_id, "employee.deactivate", "employee", employee.id, {})
    session.commit()
    return employee


def set_own_consent(session: Session, employee_id: UUID, value: bool) -> Employee:
    employee = _get(session, employee_id)
    employee.public_greeting_consent = value
    employee.updated_at = utcnow()
    write_audit(session, None, "employee.consent_self", "employee", employee.id, {"public_greeting_consent": value})
    session.commit()
    return employee


def set_own_notifications(session: Session, employee_id: UUID, announcements: bool | None, collections: bool | None) -> Employee:
    employee = _get(session, employee_id)
    current = dict(employee.notification_settings or {})
    if announcements is not None:
        current["announcements"] = announcements
    if collections is not None:
        current["collections"] = collections
    employee.notification_settings = current
    employee.updated_at = utcnow()
    session.commit()
    return employee


def bind_telegram(session: Session, invite_token: str, telegram_id: int, username: str | None) -> Employee:
    employee = session.scalar(select(Employee).where(Employee.invite_token == invite_token.strip()))
    if employee is None:
        raise DomainError("Ссылка-приглашение недействительна", "invite")
    if employee.status != "active":
        raise DomainError("Профиль сотрудника неактивен", "inactive")
    other = session.scalar(select(Employee).where(Employee.telegram_id == telegram_id, Employee.id != employee.id))
    if other is not None:
        raise DomainError("Этот Telegram уже привязан к другому сотруднику", "telegram_taken")
    employee.telegram_id = telegram_id
    if username:
        employee.telegram_username = normalize_username(username)
    employee.bot_status = "active"
    employee.updated_at = utcnow()
    write_audit(session, None, "employee.bind_telegram", "employee", employee.id, {})
    session.commit()
    return employee


def mark_bot_blocked(session: Session, employee_id: UUID) -> None:
    employee = session.get(Employee, employee_id)
    if employee is None:
        return
    employee.bot_status = "blocked"
    employee.updated_at = utcnow()
    session.commit()


def import_rows(session: Session, rows: list[dict], actor_id: UUID | None) -> dict:
    created = 0
    duplicates: list[dict] = []
    errors: list[dict] = []
    accepted: list[dict] = []
    for index, row in enumerate(rows, start=1):
        try:
            data = _row_to_input(session, row)
        except (DomainError, ValueError) as exc:
            errors.append({"row": index, "message": str(exc)})
            continue
        identity = _identity(data, source="file")
        file_matches = duplicate_reasons(identity, accepted)
        db_matches = _matches(session, data, exclude_id=None, source="database")
        if file_matches or db_matches:
            duplicates.append({"row": index, "full_name": data.full_name, "matches": file_matches + db_matches})
            continue
        employee = _build(data)
        session.add(employee)
        session.flush()
        accepted.append(_identity(data, source="file", entity_id=str(employee.id)))
        created += 1
    write_audit(
        session,
        actor_id,
        "employee.import",
        "employee",
        None,
        {"created": created, "duplicates": len(duplicates), "errors": len(errors)},
    )
    session.commit()
    return {"created": created, "duplicates": duplicates, "errors": errors}


def parse_csv(content: bytes) -> list[dict]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def parse_xlsx(content: bytes) -> list[dict]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    parsed: list[dict] = []
    for raw in rows[1:]:
        if raw is None or all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        parsed.append({headers[i]: raw[i] for i in range(len(headers)) if headers[i]})
    return parsed


def list_employees(
    session: Session,
    *,
    query: str | None = None,
    status: str | None = None,
    department_id: UUID | None = None,
    bot_status: str | None = None,
    limit: int = 100,
    offset: int = 0,
) -> list[Employee]:
    stmt = select(Employee).order_by(Employee.full_name)
    if status:
        stmt = stmt.where(Employee.status == status)
    if department_id:
        stmt = stmt.where(Employee.department_id == department_id)
    if bot_status:
        stmt = stmt.where(Employee.bot_status == bot_status)
    if query:
        like = f"%{query.strip()}%"
        stmt = stmt.where(
            or_(
                Employee.full_name.ilike(like),
                Employee.preferred_name.ilike(like),
                Employee.position.ilike(like),
                Employee.telegram_username.ilike(like),
            )
        )
    return list(session.scalars(stmt.limit(limit).offset(offset)))


def get_employee(session: Session, employee_id: UUID) -> Employee:
    return _get(session, employee_id)


def find_by_telegram(session: Session, telegram_id: int) -> Employee | None:
    return session.scalar(select(Employee).where(Employee.telegram_id == telegram_id))


def ensure_department(session: Session, name: str) -> Department:
    clean = " ".join(name.strip().split())
    if not clean:
        raise DomainError("Пустое подразделение")
    found = session.scalar(select(Department).where(Department.name == clean))
    if found:
        return found
    department = Department(name=clean, is_active=True)
    session.add(department)
    session.flush()
    return department


def _get(session: Session, employee_id: UUID) -> Employee:
    employee = session.get(Employee, employee_id)
    if employee is None:
        raise NotFoundError("Сотрудник не найден")
    return employee


def _bot_status(data: EmployeeInput) -> str:
    if data.bot_status in {"not_started", "active", "blocked"}:
        return data.bot_status
    if data.telegram_id:
        return "active"
    return "not_started"


def _validate(data: EmployeeInput) -> None:
    if not data.full_name or not data.full_name.strip():
        raise DomainError("Укажите ФИО")
    if not data.preferred_name or not data.preferred_name.strip():
        raise DomainError("Укажите имя для обращения")
    if not isinstance(data.birth_date, date):
        raise DomainError("Укажите дату рождения")


def _build(data: EmployeeInput) -> Employee:
    return Employee(
        full_name=data.full_name.strip(),
        preferred_name=data.preferred_name.strip(),
        birth_date=data.birth_date,
        department_id=data.department_id,
        position=(data.position or "").strip(),
        phone=(data.phone or "").strip(),
        city=(data.city or "").strip(),
        telegram_id=data.telegram_id,
        telegram_username=normalize_username(data.telegram_username),
        hired_at=data.hired_at,
        status=data.status if data.status in {"active", "inactive"} else "active",
        public_greeting_consent=data.public_greeting_consent,
        participate_in_collections=data.participate_in_collections,
        notification_settings=data.notification_settings
        or {"announcements": True, "collections": True},
        invite_token=secrets.token_urlsafe(24),
        bot_status=_bot_status(data),
        created_at=utcnow(),
        updated_at=utcnow(),
    )


def _identity(data: EmployeeInput, source: str, entity_id: str | None = None) -> dict:
    return {
        "id": entity_id,
        "full_name": data.full_name,
        "birth_date": data.birth_date,
        "telegram_id": data.telegram_id,
        "telegram_username": data.telegram_username,
        "source": source,
    }


def _matches(session: Session, data: EmployeeInput, exclude_id: UUID | None, source: str) -> list[dict]:
    existing = [_employee_identity(row, source) for row in session.scalars(select(Employee))]
    if exclude_id is not None:
        existing = [row for row in existing if row["id"] != str(exclude_id)]
    return duplicate_reasons(_identity(data, source), existing)


def _employee_identity(employee: Employee, source: str) -> dict:
    return {
        "id": str(employee.id),
        "full_name": employee.full_name,
        "birth_date": employee.birth_date,
        "telegram_id": employee.telegram_id,
        "telegram_username": employee.telegram_username,
        "source": source,
    }


def _row_to_input(session: Session, row: dict) -> EmployeeInput:
    normalized = {str(key).strip().lower(): value for key, value in row.items() if key is not None}
    full_name = _text(normalized.get("full_name") or normalized.get("фио"))
    preferred = _text(normalized.get("preferred_name") or normalized.get("имя") or full_name)
    birth = _date(normalized.get("birth_date") or normalized.get("дата_рождения"))
    if birth is None:
        raise DomainError("Некорректная дата рождения")
    department_name = _text(normalized.get("department") or normalized.get("подразделение"))
    department_id = ensure_department(session, department_name).id if department_name else None
    telegram_id = _optional_int(normalized.get("telegram_id"))
    consent_raw = normalized.get("public_greeting_consent", normalized.get("согласие", False))
    return EmployeeInput(
        full_name=full_name,
        preferred_name=preferred,
        birth_date=birth,
        department_id=department_id,
        position=_text(normalized.get("position") or normalized.get("должность")),
        telegram_id=telegram_id,
        telegram_username=_text(normalized.get("telegram_username") or normalized.get("username")) or None,
        hired_at=_date(normalized.get("hired_at") or normalized.get("дата_трудоустройства")),
        public_greeting_consent=_bool(consent_raw),
        participate_in_collections=_bool(normalized.get("participate_in_collections", True)),
    )


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _optional_int(value: object) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    return int(str(value).strip())


def _bool(value: object) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in TRUE_VALUES


def _date(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise DomainError(f"Некорректная дата: {text}")
